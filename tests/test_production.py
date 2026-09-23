import argparse
import json
import tempfile
import threading
import time
import unittest
from pathlib import Path
from unittest import mock
import httpx
from openpyxl import load_workbook
from fpt_einvoice.api import fetch_invoices
from fpt_einvoice.constants import UI_EXPORT_COLUMNS
from fpt_einvoice.export import EXCEL_MAX_DATA_ROWS, export_raw_workbook
from fpt_einvoice.production import run_export_year, run_merge_year, split_month

def response(rows):
    return httpx.Response(200, json={"data": rows}, request=httpx.Request("GET", "https://x/api/sea"))

class ConcurrentClient:
    def __init__(self): self.active = 0; self.maximum = 0; self.lock = threading.Lock(); self.starts = []
    def get(self, path, params):
        with self.lock:
            self.active += 1; self.maximum = max(self.maximum, self.active); self.starts.append(params["start"])
        time.sleep(.01)
        with self.lock: self.active -= 1
        return response([{"inc": params["start"]}]) if params["start"] < 5 else response([])

class ProductionTests(unittest.TestCase):
    def test_month_boundaries(self):
        for year, month, last in [(2023,1,31),(2023,2,28),(2024,2,29),(2023,4,30)]:
            ranges = split_month(year, month)
            self.assertEqual([(a.day,b.day) for a,b,_ in ranges], [(1,10),(11,20),(21,last)])
            days = [d for a,b,_ in ranges for d in range(a.day,b.day+1)]
            self.assertEqual(days, list(range(1,last+1)))

    def test_jsonl_pages_resume_and_worker_bound(self):
        with tempfile.TemporaryDirectory() as tmp:
            raw = Path(tmp) / "x.jsonl"; client = ConcurrentClient()
            rows = fetch_invoices(client, "01GTKT", "a", "b", 2, 1, workers=3, raw_path=raw)
            self.assertEqual(len(rows), 5); self.assertLessEqual(client.maximum, 3)
            before = len(client.starts)
            rows2 = fetch_invoices(client, "01GTKT", "a", "b", 2, 1, workers=3, raw_path=raw, resume=True)
            self.assertEqual(rows2, rows); self.assertEqual(len(client.starts), before)
            manifest = json.loads((raw.parent / "x.jsonl.pages" / "manifest.json").read_text())
            self.assertTrue(manifest["complete"])

    def test_no_adaptive_page_size_retries_same_count(self):
        class C:
            def __init__(self): self.counts = []
            def get(self, path, params):
                self.counts.append(params["count"])
                r = httpx.Response(502, request=httpx.Request("GET", "https://x")); r.raise_for_status()
        client = C()
        with self.assertRaises(httpx.HTTPStatusError):
            fetch_invoices(client, "01GTKT", "a", "b", 2, 5000, max_retries=2, retry_delay=0, adaptive_page_size=False)
        self.assertEqual(client.counts, [5000, 5000, 5000])

    def test_range_retry(self):
        args = argparse.Namespace(year=2023, output_dir="unused", types="all-known", range_retries=3,
            retry_delay=0, workers=5, page_size=5000, resume=True, adaptive_page_size=False)
        calls = 0
        def fake_export(_):
            nonlocal calls; calls += 1
            if calls == 1: raise RuntimeError("transient")
            return {"ok": True}
        with tempfile.TemporaryDirectory() as tmp:
            args.output_dir = tmp
            with mock.patch("fpt_einvoice.cli.run_export", side_effect=fake_export), \
                 mock.patch("fpt_einvoice.production._merge_month_raw", return_value=([], {c:0 for c in ("01GTKT","03XKNB","01/MTT","06HDTM")})), \
                 mock.patch("fpt_einvoice.production.export_raw_workbook"):
                run_export_year(args)
        self.assertEqual(calls, 37)

    def test_partial_tmp_is_not_complete(self):
        with tempfile.TemporaryDirectory() as tmp:
            raw = Path(tmp) / "x.jsonl"; pages = raw.parent / "x.jsonl.pages"; pages.mkdir()
            (pages / "000000000.jsonl.tmp").write_text('{"inc":1}\n')
            (pages / "manifest.json").write_text(json.dumps({"version":1,"type":"01GTKT","fd":"a","td":"b","page_size":2,"complete":False}))
            class C:
                def get(self, p, params): return response([])
            fetch_invoices(C(), "01GTKT", "a", "b", 2, 2, raw_path=raw, resume=True)
            self.assertEqual(raw.read_text(), "")

    def test_fast_excel_schema_and_no_duplicate_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp); raw = root / "x.jsonl"; raw.write_text('{"inc":7,"idt":"2023-01-01"}\n')
            out = root / "x.xlsx"; export_raw_workbook(out, [("01GTKT",raw)], {"types":["01GTKT"],"counts":{"01GTKT":1},"total_rows":1})
            wb = load_workbook(out, read_only=True)
            self.assertEqual(wb.sheetnames, ["metadata","summary","invoices_all"])
            self.assertEqual(list(next(wb["invoices_all"].iter_rows(values_only=True))), [h for _,h in UI_EXPORT_COLUMNS])
            self.assertEqual(sum(1 for _ in wb["invoices_all"].iter_rows(values_only=True)), 2)

    def test_row_limit_detection(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "row limit"):
                export_raw_workbook(Path(tmp)/"x.xlsx", [], {"total_rows":EXCEL_MAX_DATA_ROWS+1})

    def test_merge_year_raw_fast_path_and_xlsx_fallback(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for month in range(1,13):
                key = f"2023-{month:02d}"; md = root/key; (md/"raw").mkdir(parents=True)
                raw = md/"raw"/"01GTKT.jsonl"; raw.write_text("" if month == 2 else json.dumps({"inc":month})+"\n")
            # Force January fallback by removing raw after making monthly XLSX.
            jan = root/"2023-01"; export_raw_workbook(jan/"fpt_einvoice_2023-01.xlsx", [("01GTKT",jan/"raw"/"01GTKT.jsonl")], {"types":["01GTKT"],"counts":{"01GTKT":1},"total_rows":1})
            (jan/"raw"/"01GTKT.jsonl").unlink()
            result = run_merge_year(argparse.Namespace(year=2023, output_dir=str(root), skip_missing=False))
            self.assertEqual(result["sources"]["2023-01"], "xlsx"); self.assertEqual(result["sources"]["2023-02"], "jsonl")
            wb = load_workbook(result["output_xlsx"], read_only=True)
            self.assertEqual(wb.sheetnames, [f"2023-{m:02d}" for m in range(1,13)])
            self.assertEqual(sum(1 for _ in wb["2023-02"].iter_rows(values_only=True)), 1)

if __name__ == "__main__": unittest.main()
