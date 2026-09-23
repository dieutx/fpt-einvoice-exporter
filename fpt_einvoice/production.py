"""Year/month orchestration and streaming yearly workbook merge."""
import calendar
import json
import os
import time
from argparse import Namespace
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator
from openpyxl import Workbook, load_workbook
from .constants import KNOWN_TYPE_LABELS, UI_EXPORT_COLUMNS
from .export import EXCEL_MAX_DATA_ROWS, build_ui_export_row, export_raw_workbook, iter_checkpoint_rows, write_json, _append_write_only_header
from .log import eprint

def split_month(year: int, month: int) -> list[tuple[date, date, str]]:
    last = calendar.monthrange(year, month)[1]
    return [(date(year, month, 1), date(year, month, 10), "01_10"),
            (date(year, month, 11), date(year, month, 20), "11_20"),
            (date(year, month, 21), date(year, month, last), "21_END")]

def _safe_type(code: str) -> str: return code.replace("/", "_")

def _merge_month_raw(month_dir: Path, ranges: list[tuple[date, date, str]], types: list[str]) -> tuple[list[tuple[str, Path]], dict[str, int]]:
    target_dir = month_dir / "raw"; target_dir.mkdir(parents=True, exist_ok=True)
    outputs = []; counts = {}
    for code in types:
        destination = target_dir / f"{_safe_type(code)}.jsonl"; tmp = destination.with_name(destination.name + ".tmp")
        count = 0
        with tmp.open("w", encoding="utf-8") as out:
            for _, _, label in ranges:
                source = month_dir / "parts" / label / "raw" / f"{_safe_type(code)}.jsonl"
                if not source.exists(): continue
                for row in iter_checkpoint_rows(source):
                    # Date ranges never overlap and each page is assembled once from
                    # its manifest, so streaming concatenation cannot create resume duplicates.
                    out.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n"); count += 1
            out.flush(); os.fsync(out.fileno())
        os.replace(tmp, destination); outputs.append((code, destination)); counts[code] = count
        eprint(f"[merge-raw] {code} rows={count}")
    return outputs, counts

def run_export_year(args: Namespace) -> dict[str, Any]:
    from .cli import run_export
    root = Path(args.output_dir).expanduser().resolve(); results = []
    types = list(KNOWN_TYPE_LABELS) if args.types == "all-known" else [x.strip() for x in args.types.split(",") if x.strip()]
    for month in range(1, 13):
        month_key = f"{args.year:04d}-{month:02d}"; month_dir = root / month_key; eprint(f"[month] {month_key}")
        ranges = split_month(args.year, month)
        for start, end, label in ranges:
            part_dir = month_dir / "parts" / label; eprint(f"[range] {start.isoformat()} -> {end.isoformat()}")
            error = None
            for attempt in range(1, args.range_retries + 1):
                part_args = Namespace(**vars(args)); part_args.from_date = start.isoformat(); part_args.to_date = end.isoformat()
                part_args.output_dir = str(part_dir); part_args.output_name = None; part_args.raw_only = True
                part_args.adaptive_page_size = False; part_args.continue_on_error = False
                part_args.resume = args.resume or attempt > 1
                try:
                    run_export(part_args); error = None; eprint(f"[range-ok] {label}"); break
                except (KeyboardInterrupt, SystemExit): raise
                except Exception as exc:
                    error = exc
                    if attempt < args.range_retries:
                        eprint(f"[range-retry] attempt={attempt + 1}/{args.range_retries} error={exc}"); time.sleep(args.retry_delay)
            if error is not None: raise RuntimeError(f"Range {start} -> {end} thất bại sau {args.range_retries} lần") from error
        raw_paths, counts = _merge_month_raw(month_dir, ranges, types); total = sum(counts.values())
        metadata = {"year_month": month_key, "types": types, "counts": counts, "total_rows": total,
                    "generated_at": datetime.now().astimezone().isoformat(), "workers": args.workers, "page_size": args.page_size}
        write_json(month_dir / "metadata.json", metadata)
        if total == 0: eprint(f"[empty-month] {month_key}")
        output = month_dir / f"fpt_einvoice_{month_key}.xlsx"; export_raw_workbook(output, raw_paths, metadata)
        eprint(f"[month-ok] {month_key} rows={total}"); results.append(str(output))
    return {"ok": True, "year": args.year, "months": results}

def _month_raw_paths(month_dir: Path) -> list[Path]:
    raw = month_dir / "raw"
    if not raw.exists(): return []
    paths = []
    for code in KNOWN_TYPE_LABELS:
        path = raw / f"{_safe_type(code)}.jsonl"
        if path.exists(): paths.append(path)
    return paths

def run_merge_year(args: Namespace) -> dict[str, Any]:
    root = Path(args.output_dir).expanduser().resolve(); output = root / f"fpt_einvoice_{args.year}_all_months.xlsx"
    wb = Workbook(write_only=True); headers = [h for _, h in UI_EXPORT_COLUMNS]; sources = {}
    for month in range(1, 13):
        key = f"{args.year:04d}-{month:02d}"; month_dir = root / key; raw_paths = _month_raw_paths(month_dir)
        xlsx = month_dir / f"fpt_einvoice_{key}.xlsx"
        if not raw_paths and not xlsx.exists():
            if args.skip_missing: eprint(f"[skip-missing] {key}"); continue
            raise FileNotFoundError(f"Chưa export tháng {key}: {xlsx}")
        ws = wb.create_sheet(key); _append_write_only_header(ws, headers); count = 0
        if raw_paths:
            sources[key] = "jsonl"
            for path in raw_paths:
                for raw in iter_checkpoint_rows(path):
                    if count >= EXCEL_MAX_DATA_ROWS: raise ValueError(f"Excel row limit exceeded for {key}: > {EXCEL_MAX_DATA_ROWS}")
                    row = build_ui_export_row(raw); ws.append([row.get(h, "") for h in headers]); count += 1
        else:
            sources[key] = "xlsx"; source_wb = load_workbook(xlsx, read_only=True, data_only=True)
            try:
                source_ws = source_wb["invoices_all"]
                first = True
                for values in source_ws.iter_rows(values_only=True):
                    if first: first = False; continue
                    if count >= EXCEL_MAX_DATA_ROWS: raise ValueError(f"Excel row limit exceeded for {key}: > {EXCEL_MAX_DATA_ROWS}")
                    ws.append(list(values)); count += 1
            finally: source_wb.close()
        eprint(f"[merge-year] {key} source={sources[key]} rows={count}")
    output.parent.mkdir(parents=True, exist_ok=True); wb.save(output)
    return {"ok": True, "output_xlsx": str(output), "sources": sources}
